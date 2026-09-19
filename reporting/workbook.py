"""The workbook export: the comparison table's two stacked blocks as a spreadsheet.

The workbook is the same content as the print, written from the same structure, which is why
`compare/table.py` builds the blocks once and both readers render them: a workbook assembled by hand
from the same numbers would be a second copy of the table that drifts the first time a column moves.

**Every sheet carries the provenance block.** The snapshot, the protocol, the cost multiple and the
bar are written at the top of each sheet rather than on a cover page, because a sheet is what gets
copied out, and a number without its snapshot and its cost multiple cannot be checked by whoever
receives it. The cell id is the first column of every row for the same reason.

**Demonstration depth, and it says so.** This exports the table, which is the deliverable a reader
opens Excel for; it does not emulate a reporting platform, chart the cells, or export any other layer.

Nothing in the analytics imports this module, and the acceptance fixture checks the direction over the
import graph rather than by inspection.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from portfolio_workbench.compare import table as table_module
from portfolio_workbench.data import loader

HERE = Path(__file__).resolve().parent
DEFAULT_OUTPUT_ROOT = HERE.parent / ".data" / "runs"

# The workbook's own renderings of the print's format codes. A number is written as a number so the
# sheet stays sortable and summable, and the format code travels with the column that declared it.
FORMATS = {".2%": "0.00%", ".0%": "0%", ".2f": "0.00", ".3f": "0.000", "s": "General"}

# The two sheets the design fixes, in the order the blocks are read. Excel refuses a sheet name past
# thirty-one characters, which these are not.
SHEET_NAMES = {
    table_module.RETURNS_BLOCK: "returns and risk",
    table_module.COST_BLOCK: "cost and diagnostics",
}

BOLD = Font(bold=True)


def write(sheet, document, path):
    """Write the two blocks, each with the provenance block above its tables, and return the path."""
    workbook = Workbook()
    provenance = table_module.provenance(sheet, document)
    for position, block in enumerate(table_module.block_tables(sheet)):
        if position == 0:
            worksheet = workbook.active
            worksheet.title = SHEET_NAMES[block["block"]]
        else:
            worksheet = workbook.create_sheet(SHEET_NAMES[block["block"]])
        _write_provenance(worksheet, provenance)
        widths = {}
        for table in block["tables"]:
            _write_table(worksheet, table, widths)
        for note in block["notes"]:
            worksheet.append([note])
        _set_widths(worksheet, widths)
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return target


def _write_provenance(worksheet, provenance):
    """The block every sheet is read under, repeated at the top of each sheet rather than on a cover."""
    for name, value in provenance.items():
        worksheet.append([name, value])
        worksheet.cell(row=worksheet.max_row, column=1).font = BOLD
    worksheet.append([])


def _write_table(worksheet, table, widths):
    """One table: its title, its column names and its rows, each column under its own number format.

    A column width is accumulated across the tables of one sheet rather than set per table, because
    Excel sizes a column for the whole sheet and the last table written would otherwise decide the
    width of a column an earlier table needs wider.
    """
    worksheet.append([table["title"]])
    worksheet.cell(row=worksheet.max_row, column=1).font = BOLD
    worksheet.append([name for name, width, spec, align in table["columns"]])
    first = worksheet.max_row + 1
    for row in table["rows"]:
        worksheet.append([None if value is None else value for value in row])
    for position, (name, width, spec, align) in enumerate(table["columns"], start=1):
        for offset in range(len(table["rows"])):
            worksheet.cell(row=first + offset, column=position).number_format = FORMATS.get(spec, "General")
        widths[position] = max(widths.get(position, 0), width, len(str(name)))
    worksheet.append([])


def _set_widths(worksheet, widths):
    """The column widths, wide enough for the widest name or value their column was declared with."""
    for position, width in widths.items():
        worksheet.column_dimensions[worksheet.cell(row=1, column=position).column_letter].width = width + 2


def main(document=None, path=None):
    """Analyse the snapshot and write the workbook beside the run manifests."""
    from portfolio_workbench import study

    analysis = study.analyse(loader.load_panel() if document is None else document)
    if path is None:
        path = DEFAULT_OUTPUT_ROOT / analysis.document.snapshot_id / "comparison.xlsx"
    target = write(analysis.sheet, analysis.document, path)
    print(
        f"[table] workbook written to {target}: {len(table_module.block_tables(analysis.sheet))} blocks, each "
        f"sheet carrying the snapshot {analysis.document.snapshot_id}, the protocol and the cost multiple"
    )
    return target


if __name__ == "__main__":
    main()
