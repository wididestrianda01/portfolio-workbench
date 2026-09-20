"""One end-to-end run on the frozen snapshot, asserting the structure the whole build rests on.

This is the fixture the rest of the suite extends: the structural checks a reader is asked to take on
trust are asserted here rather than left in prose, so a change that breaks one of them fails one
command. It runs the grid, the comparison table, the attribution and the budget once, at module scope,
through the analysis, and reads every claim from that one run. Two tests then make a pass of their own,
because that pass **is** what they check: the rerun test compares a second grid against the first, and
the check on the prose deliverables re-reads the code the published documents were written from.

**The rerun tolerance is stated, not implied.** A rerun of the grid on the frozen snapshot reproduces
the metric table within `statistics.RERUN_TOLERANCE` (1e-06 relative), and every stochastic step draws
from the documented seed (`statistics.SEED`, `cp.SEED`), so two runs of this fixture are the same run.
The one deliberate inequality is the cost sensitivity: a higher per-side multiple may only lower an
information ratio, which is asserted as a direction rather than as a value.
"""

import ast
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from portfolio_workbench import study
from portfolio_workbench.attribute import brinson
from portfolio_workbench.budget import euler
from portfolio_workbench.compare import grid as grid_module
from portfolio_workbench.compare import registry, table as table_module
from portfolio_workbench.construct import constraints, families, means
from portfolio_workbench.data import loader, sql, universe
from portfolio_workbench.evaluate import metrics, statistics, walkforward
from portfolio_workbench.factors import components as cp
from portfolio_workbench.factors import exposures, spanning
from portfolio_workbench.factors import spine as spine_module
from portfolio_workbench.risk import covariance
from reporting import skills_matrix, source_map, workbook


# The names a run's manifest and its table row share. They are one vocabulary now, which is what lets a
# manifest be joined to the row it belongs to rather than read through the function that renamed it.
MANIFEST_ROW_NAMES = (
    "turnover_annualised",
    "cost_annualised",
    "cap_binding_frequency",
    "weight_stability",
    "concentration",
    "largest_weight",
)


@pytest.fixture(scope="module")
def run():
    document = loader.load_panel()
    months = document.months
    returns = document.returns
    block = spine_module.constructed_block(returns)
    named = spine_module.named_set(document.factors.eur, block)
    return {
        "document": document,
        "months": months,
        "returns": returns,
        "block": block,
        "named": named,
        "fit": exposures.rolling(returns, named, months),
    }


def test_the_factor_layer_runs_end_to_end_on_the_frozen_snapshot(run):
    returns, named, months = run["returns"], run["named"], run["months"]
    fit = run["fit"]

    # The data contract: the joined panel that every result is keyed to.
    assert len(months) == 191 and str(months.min()) == "2010-09" and str(months.max()) == "2026-07"
    assert list(returns.columns) == universe.TICKERS
    assert not returns.isna().to_numpy().any()

    # The named set: the published spine, the constructed block, no gap, and the file's own risk-free
    # refused as a factor.
    assert "RF" not in named.columns
    assert len(named) == 190 and len(named.columns) == 10
    assert not named.isna().to_numpy().any()

    # The walk-forward: the decided out-of-sample calendar, one observation short only in the first
    # window, and every sleeve fitted on the design it admits.
    traded = fit["traded"]
    assert len(traded) == 131 and traded[0] == pd.Period("2015-09", freq="M")
    assert traded[-1] == pd.Period("2026-07", freq="M")
    assert fit["n_obs"][0] == exposures.MIN_OBS and set(fit["n_obs"][1:]) == {exposures.WINDOW}

    # The four constructed sleeves are the block, so their block loadings are identities and their
    # model is exact; the other seven carry an alpha with a t-statistic and a share of variance.
    identity = fit["identity"][-1]
    expected = {
        "IBGL.AS": [1.0, 0.5, 0.0, 0.0],
        "IEGE.AS": [1.0, -0.5, 0.0, 0.0],
        "IEAC.AS": [1.0, 0.5, 1.0, 0.0],
        "IHYG.L": [1.0, 0.5, 1.0, 1.0],
    }
    assert set(identity) == set(expected)
    for sleeve, loadings in identity.items():
        assert loadings == pytest.approx(expected[sleeve], abs=1e-9), f"{sleeve}'s construction changed"
    assert fit["r2"]["IBGL.AS"].min() == pytest.approx(1.0)
    assert fit["resid_var"]["IBGL.AS"].max() == pytest.approx(0.0, abs=1e-24)
    assert fit["alpha"]["IBGL.AS"].isna().all(), "no alpha to test: the block is the construction"
    estimated = [sleeve for sleeve in returns.columns if sleeve not in identity]
    assert len(estimated) == 7
    assert fit["alpha"][estimated].notna().to_numpy().all()
    assert (fit["r2"][estimated] < 1.0).to_numpy().all()

    # The count rule: inside the pre-registered bound, stable across permutation seeds, and not
    # falsified by its own movement.
    rule = cp.decompose(returns)
    assert 1 <= rule["components"] <= cp.PREREGISTERED_K
    series = cp.count_series(returns, months)
    assert series["falsified"] is False
    assert set(series["counts"]) <= {1, 2, 3}
    assert cp.stability(returns, months)["agrees"] is True

    # Spanning: the verdicts at the count the rule retains, in both directions.
    directions = spanning.directions(returns, named, rule["components"])
    assert directions["headline"]["p_value"] < 0.05 and directions["reverse"]["p_value"] < 0.05
    assert np.abs(directions["headline"]["rows_sum"] - 1.0).max() < 1.0

    # The covariance axis: three estimators on one window, each labelled, symmetric and positive
    # semi-definite, with the shrinkage estimator conditioning the sample's problem away.
    window = months[-60:]
    block = exposures.window_block(returns, window)
    report = covariance.conditioning(block)
    assert report["sample_condition"] > report["shrinkage_condition"]
    assert 0.0 < report["intensity"] < 1.0
    for name, matrix in report["covariances"].items():
        assert list(matrix.index) == universe.TICKERS, f"{name} lost the sleeve map's order"
        assert np.allclose(matrix.to_numpy(), matrix.to_numpy().T)
        assert np.all(np.linalg.eigvalsh(matrix.to_numpy()) > -1e-18), f"{name} is not positive semi-definite"


def test_nothing_in_the_analytics_imports_the_reporting_layer():
    """The dependency direction the package shape fixes, checked over the import graph.

    The rule is not stylistic: the analytics have to run on a machine with no reporting package and no
    spreadsheet library, and an import of the reporter from inside a module would make the workbook's
    dependencies the engine's. Reading the source for the string would pass on a commented-out line and
    fail on an aliased import, so the check parses each module and reads the import statements.
    """
    root = Path(__file__).resolve().parents[1] / "portfolio_workbench"
    modules = sorted(root.rglob("*.py"))
    assert modules, f"no analytics modules found under {root}"
    offenders = []
    for path in modules:
        for node in ast.walk(ast.parse(path.read_text())):
            if isinstance(node, ast.Import):
                names = [alias.name for alias in node.names]
            elif isinstance(node, ast.ImportFrom):
                names = [node.module or ""]
            else:
                continue
            if any(name.split(".")[0] == "reporting" for name in names):
                offenders.append(f"{path.relative_to(root)}:{node.lineno}")
    assert not offenders, f"the analytics imports the reporting layer at {offenders}"


def test_the_layer_entry_points_reach_the_analysis_only_inside_their_functions():
    """The one edge that runs against the layer order, and the rule that keeps it from becoming a cycle.

    `study` sits above the layers and imports them, so a layer module that imported it at module level
    would be a cycle: Python would refuse it outright, or accept it in one import order and not the
    other, which is worse. Every entry point therefore reaches the analysis inside its own `main`, and
    this check holds that deferral in place rather than leaving it to a comment. The edge is deliberate
    and documented in the module's own header; an import moved to module level is a change to the
    package's shape rather than a tidy-up.
    """
    root = Path(__file__).resolve().parents[1] / "portfolio_workbench"
    deferred = []
    for path in sorted(root.rglob("*.py")):
        if path.name == "study.py":
            continue
        tree = ast.parse(path.read_text())
        for node in ast.walk(tree):
            if isinstance(node, ast.ImportFrom):
                module = (node.module or "").split(".")[-1]
                names = [alias.name.split(".")[-1] for alias in node.names]
            elif isinstance(node, ast.Import):
                module, names = "", [alias.name.split(".")[-1] for alias in node.names]
            else:
                continue
            if module != "study" and "study" not in names:
                continue
            assert node not in tree.body, (
                f"{path.relative_to(root)}:{node.lineno} imports the analysis at module level, which the "
                "layer order cannot accept: the analysis imports this module"
            )
            deferred.append(f"{path.relative_to(root)}:{node.lineno}")
    assert deferred, "nothing reaches the analysis, so this check is measuring nothing"


def test_the_sql_statement_of_the_contract_agrees_with_the_pandas_path():
    """The table contract is stated twice - as pandas code the analytics run on, and as a query - and
    the two statements are checked against each other on the frozen snapshot.

    The check is not about the query's speed or elegance. The contract is what another project has to
    satisfy to consume this engine, and a contract written only as Python is one the sibling re-derives
    instead of satisfying. Two statements that disagree mean one of them is wrong about the panel every
    result is keyed to, so the disagreement is refused rather than noted.

    The same holds for the table itself, not only for its columns: the monthly euro excess return frame
    the factor model is fitted on is assembled in SQL from the snapshot's own files - currency
    translation, cash accrual and gate included - and compared against the pandas path cell by cell.
    """
    document = loader.load_panel()
    con, _, views = sql.connection()
    checked = sql.agreement(document, None)
    assert checked["agrees"], checked["differences"]

    # The gate is the rule, so it is asserted where it bites: read at the last joined month rather than
    # at the snapshot's own date, the query hides the bars that month could not have seen.
    at_snapshot = sql.coverage(con, views, when=document.as_of)
    earlier = sql.coverage(con, views, when=document.months.max().to_timestamp())
    assert earlier["months_visible"].sum() < at_snapshot["months_visible"].sum()
    assert earlier["months_visible"].sum() < earlier["months"].sum()

    assembled = sql.returns_agreement(document, sql.returns(con, views, document))
    assert assembled["agrees"], assembled["differences"]
    assert (assembled["instruments"], assembled["months"]) == (11, 191)
    assert assembled["worst"] <= sql.RETURN_TOLERANCE


@pytest.fixture(scope="module")
def stack():
    """One run of the whole stack on the frozen snapshot: the analysis every reading below is taken from.

    Assembled once here through the analysis rather than by hand, because the hand-built version is what
    let two readings of one run be taken against two different covariances.
    """
    analysis = study.analyse(loader.load_panel())
    return {
        "document": analysis.document,
        "grid": analysis.grid,
        "benchmark": analysis.benchmark,
        "sheet": analysis.sheet,
        "policy": analysis.policy,
        "covariance": analysis.covariance,
        "holding": list(analysis.attribution.values()),
        "analysis": analysis,
    }


def test_the_analysis_reads_one_covariance_over_the_traded_months(stack):
    """Every reading of the risk budget is taken against one covariance, over one window.

    The defect this guards is not a wrong number but two right-looking ones: the budget was read against
    the panel's own returns frame at one call site and against the traded months at another, so two
    published documents stated different contributions for the same book and nothing compared them. The
    check pins the months the covariance is taken over, and pins it to the estimator module's own sample
    covariance rather than a fourth copy of the arithmetic in the layer that reports it.
    """
    analysis = stack["analysis"]
    traded = analysis.grid["results"][0]["traded"]
    assert analysis.covariance.equals(covariance.sample(analysis.returns.reindex(traded)))
    # The panel's own returns frame is the window this was read over before, and it is not this one: a
    # budget read over it describes a book held for months the run never held it.
    assert not analysis.covariance.equals(covariance.sample(analysis.returns.iloc[1:]))


def test_a_run_manifest_and_its_table_row_carry_one_set_of_names(stack, tmp_path):
    """A run's manifest is what makes its weight path reproducible, and nothing read one back.

    Its leaves agree with the table row's because the runner and the row now use one vocabulary: the row
    used to rename three of the runner's measurements, so joining a manifest to a row meant reading the
    function that did the renaming, and the manifest's own names could drift from the code unchecked.
    The check joins them on the cell id - the name a manifest is written under - and holds the shared
    names against the row, which is what makes the record checkable rather than merely written.
    """
    written = grid_module.write_manifests(stack["document"], stack["grid"]["results"], tmp_path)
    by_cell = {row["cell"]: row for row in stack["sheet"]["rows"]}
    assert len(written) == len(stack["grid"]["results"])
    for path in written:
        record = json.loads(path.read_text())
        row = by_cell[record["cell"]]
        for name in MANIFEST_ROW_NAMES:
            assert record["summary"][name] == pytest.approx(row[name]), f"{record['cell']}: {name}"


def test_the_cell_count_is_the_pre_registered_one_and_every_cell_carries_a_verdict(stack):
    """Done criterion 7: negative results are published rather than dropped.

    The reported count is the pre-registered count as arithmetic - sixteen cells, two perturbations,
    two repeats - and every row carries a verdict and a resolution limit beside it, including the rows
    whose verdict is that no difference was detected.
    """
    sheet = stack["sheet"]
    assert len(sheet["cells"]) == len(registry.CELLS) == 16
    assert len(sheet["rows"]) == registry.PRE_REGISTERED == 20
    for row in sheet["rows"]:
        assert row["verdict"], f"{row['cell']} carries no verdict"
        assert row["paired_benchmark"]["resolution"] is not None
    negative = table_module.negative_results(sheet)
    assert negative, "the run published no negative result, which the design says is a result"
    assert all(entry["verdict"] in table_module.NEGATIVE for entry in negative)


def test_the_cost_convention_is_the_traded_notional_one(stack):
    """Done criterion 11: the factor of two cannot return silently.

    Cost is `2 x one-way turnover x per-side rate`, charged on traded notional. The check reads the
    convention from the function that owns it and then reads the run: what the runner charged is what
    the convention says it charged, to the last decimal.
    """
    one_way = 0.05
    assert constraints.cost(one_way) == pytest.approx(2.0 * one_way * constraints.COST_BP / 10_000.0)
    for result in stack["grid"]["results"]:
        charged = float((result["gross"] - result["net"]).sum())
        expected = float(constraints.cost(result["turnover"].iloc[1:]).sum())
        assert charged == pytest.approx(expected, rel=1e-12), f"{result['id']} charged {charged}"


def test_the_turnover_cap_holds_and_every_cell_reports_its_establishment_cost(stack):
    """Done criterion 10: the cap holds on every measured rebalance, and the setup line is present.

    The establishment trade is the whole book traded once outside the window, so it is outside the cap
    by construction and reported as its own line rather than amortised into the monthly charges.
    """
    for result in stack["grid"]["results"]:
        measured = result["turnover"].iloc[1:]
        assert float(measured.max()) <= constraints.TURNOVER_CAP + 1e-12, f"{result['id']} breached the cap"
        establishment = result["summary"]["establishment"]
        assert establishment["turnover"] > 0.0
        assert establishment["cost"] == pytest.approx(2.0 * establishment["turnover"] * constraints.COST_BP / 10_000.0)


def test_the_attribution_reconciles_after_linking(stack):
    """Done criterion 1, on the frozen snapshot rather than on a hand case."""
    worst = 0.0
    for block in stack["holding"]:
        linked = block["linked"]
        assert linked["reconciles"], linked["relative"]
        worst = max(worst, linked["relative"])
        # A month's active return is of order a few percent, so the bound is absolute at 1e-12 and
        # the design's relative bar remains the linked total's, asserted below.
        assert np.allclose(
            block["lines"][list(brinson.LINES)].sum(axis=1).to_numpy(),
            block["lines"]["active"].to_numpy(),
            rtol=0.0,
            atol=1e-12,
        )
    assert worst < 1e-9


def test_the_euler_contributions_sum_to_volatility(stack):
    """Done criterion 2: the risk budget adds up, on every run rather than on one."""
    for result in stack["grid"]["results"]:
        checked = euler.additivity(result["weights"].mean(), stack["covariance"])
        assert checked["reconciles"], checked["relative"]
        assert checked["relative"] < 1e-10


def test_the_pca_reconstruction_is_bounded(stack):
    """Done criterion 3: the rank-k residual is the discarded eigenvalues, not an approximation.

    The residual is **measured**, not restated: the correlation matrix the rule decomposed is
    reconstructed from the retained components, and what the reconstruction leaves has the discarded
    eigenvalues as its own spectrum. Comparing `sum(eigenvalues[k:])` with
    `sum(eigenvalues) - sum(eigenvalues[:k])` would put the same quantity on both sides of the
    comparison - true for any vector and any k - and would pass on a rule whose components were
    wrong, which is the failure this criterion exists to catch.
    """
    returns = stack["grid"]["returns"]
    rule = cp.decompose(returns)
    eigenvalues = rule["eigenvalues"]
    kept = rule["components"]
    # The matrix the module decomposes, stated here rather than reached for: the correlation of the
    # sleeve excess returns, diagonal ones, trace equal to the sleeve count.
    correlation = np.corrcoef(returns.to_numpy(), rowvar=False)
    loadings = rule["loadings"]
    residual = correlation - loadings @ loadings.T
    discarded = eigenvalues[kept:]

    assert np.isclose(np.trace(residual), float(discarded.sum()), atol=1e-10)
    assert np.isclose(np.linalg.norm(residual, "fro"), float(np.sqrt((discarded ** 2).sum())), atol=1e-10)
    # The residual is rank-(sleeves - k) and its non-zero spectrum **is** the discarded one, with the
    # rest of the spectrum at zero. Comparing the residual's own eigenvalues with the discarded ones is
    # the statement the criterion makes; the previous form compared a sum with itself.
    spectrum = np.sort(np.linalg.eigvalsh(residual))[::-1]
    assert np.allclose(spectrum[: discarded.size], discarded, atol=1e-10), (
        "the residual's non-zero spectrum is the discarded eigenvalues"
    )
    assert np.allclose(spectrum[discarded.size:], 0.0, atol=1e-10), (
        "and the reconstruction leaves nothing else, so the residual's rank is the discarded count"
    )
    assert rule["variance_share"].sum() + float(discarded.sum()) / len(eigenvalues) == pytest.approx(1.0, abs=1e-10)


def test_the_black_litterman_hand_case_passes():
    """Done criterion 4: a two-asset, one-view posterior against its closed form.

    Hand-derived: with a diagonal covariance (0.04, 0.01), a prior of (1%, 2%), one view on the first
    asset at 3% and view variance 4e-4 at tau 1, the posterior's first mean is
    0.01 + 0.04 / (0.04 + 0.0004) * (0.03 - 0.01) = 0.02980198 and the second is unchanged at 0.02,
    because the view says nothing about it.
    """
    posterior = means.posterior(
        np.array([[0.04, 0.0], [0.0, 0.01]]),
        np.array([0.01, 0.02]),
        1.0,
        np.array([[1.0, 0.0]]),
        np.array([0.03]),
        np.array([[0.0004]]),
    )
    assert np.asarray(posterior).ravel() == pytest.approx([0.0298019801980198, 0.02], abs=1e-10)


def test_the_walk_forward_is_look_ahead_free(stack):
    """Done criterion 5: the boundary holds, and the assertion that checks it is part of the run."""
    steps = walkforward.steps(stack["document"].months)
    walkforward.assert_no_look_ahead(steps, stack["document"].months)
    for step in steps:
        assert step["window_end"] < step["traded"]
    traded = stack["grid"]["results"][0]["traded"]
    assert len(traded) == 131 and str(traded.min()) == "2015-09" and str(traded.max()) == "2026-07"


def test_a_rerun_reproduces_the_metric_table_and_the_tolerance_is_stated(stack):
    """The fixture's own reproducibility claim: same snapshot, same seed, same numbers."""
    again = grid_module.run_grid(stack["document"])
    benchmark = grid_module.benchmark(again["returns"], again["months"])
    checked = table_module.reproduces(
        table_module.metric_table(stack["sheet"]),
        table_module.metric_table(table_module.rows(again, benchmark)),
    )
    assert checked["agrees"], f"a rerun moved {checked['at']} by {checked['worst']:.2e}"
    assert checked["tolerance"] == statistics.RERUN_TOLERANCE


def test_the_notebook_set_covers_the_package_and_is_the_generated_one():
    """The notebooks are generated from the code, so the claim that they cannot drift is checkable.

    Three ways they could stop being true, each asserted rather than assumed: a module added to the
    package and driven by no notebook, a module claimed by two notebooks, and a module header or a
    source line edited without regenerating the notebook that renders it. The last is the one that
    matters - it is the failure that leaves a notebook describing code that no longer exists. Outputs
    are not compared, because they carry a kernel's own timing.
    """
    import nbformat

    from reporting import notebooks as notebooks_module

    root = Path(__file__).resolve().parents[1]
    driven = [module for record in notebooks_module.RECORDS for module in record["drives"]]
    assert len(driven) == len(set(driven)), "a module is driven by two notebooks"
    present = {
        str(path.relative_to(root / "portfolio_workbench"))
        for path in (root / "portfolio_workbench").rglob("*.py")
        if path.name != "__init__.py"
    }
    assert set(driven) == present, f"undriven {sorted(present - set(driven))}"

    for record in notebooks_module.RECORDS:
        path = root / "notebooks" / f"{record['slug']}.ipynb"
        assert path.exists(), f"{path.name} was never written"
        written = nbformat.read(path, as_version=4)
        errors = [output for cell in written.cells for output in cell.get("outputs", []) if output.get("output_type") == "error"]
        assert not errors, f"{path.name} carries {len(errors)} failed cell(s)"
        assert all(cell.get("execution_count") for cell in written.cells if cell.cell_type == "code"), f"{path.name} has an unexecuted cell"
        rendered = notebooks_module.notebook(record)
        assert [cell.source for cell in written.cells] == [cell.source for cell in rendered.cells], (
            f"{path.name} is not what the code renders now; run python3 -m reporting.notebooks"
        )


def test_the_consumer_boundary_carries_the_contract_and_the_five_entry_points():
    """The one import surface a consumer adopts, read against a declaration made outside the module.

    A boundary is only worth pinning if its contents are stated somewhere other than itself, so the
    five entry points and the modules behind them are declared here and the module is read against that
    declaration. A group dropped from the boundary, or a module dropped from a group, fails this rather
    than quietly narrowing what a consumer that pinned the revision can call; and identity rather than
    equality is asserted, because the promise the boundary makes is that these are the layer's own
    modules rather than copies of them.
    """
    import importlib

    from portfolio_workbench import facade

    declared = {
        "contract": (
            "portfolio_workbench.data.manifest",
            "portfolio_workbench.data.loader",
            "portfolio_workbench.data.panel",
            "portfolio_workbench.data.quality",
            "portfolio_workbench.data.sql",
        ),
        "factor_exposures": (
            "portfolio_workbench.factors.exposures",
            "portfolio_workbench.factors.components",
            "portfolio_workbench.factors.spanning",
            "portfolio_workbench.factors.spine",
        ),
        "risk_models": ("portfolio_workbench.risk.covariance",),
        "construction": (
            "portfolio_workbench.construct.families",
            "portfolio_workbench.construct.means",
            "portfolio_workbench.construct.constraints",
        ),
        "evaluation": (
            "portfolio_workbench.evaluate.walkforward",
            "portfolio_workbench.evaluate.metrics",
            "portfolio_workbench.evaluate.statistics",
        ),
        "attribution": (
            "portfolio_workbench.attribute.brinson",
            "portfolio_workbench.attribute.factor",
            "portfolio_workbench.budget.euler",
        ),
    }
    assert [name for name, _ in facade.GROUPS] == list(declared)
    for group_name, paths in declared.items():
        group = getattr(facade, group_name)
        bound = vars(group)
        assert sorted(bound) == sorted(path.rsplit(".", 1)[-1] for path in paths), group_name
        for path in paths:
            member = path.rsplit(".", 1)[-1]
            assert bound[member] is importlib.import_module(path), (
                f"{group_name}.{member} is not the module itself, so a consumer would be calling a copy"
            )


def test_the_prose_deliverables_are_the_ones_the_code_writes(tmp_path):
    """The memo, the record and the guide are generated, so the copies in the repository are read
    against the code.

    The drift this catches is silent, and it had happened: a sentence rewritten in the generator left
    the published document stating the old one, and no other check reads those files. The run
    `collect` makes is the one the documents are written from, so the comparison is byte for byte -
    the price is a second pass over the stack, which is what makes the published text a fact about the
    code rather than a file someone remembered to regenerate. All three documents are checked against
    the one evidence object, so the check costs one pass rather than three.

    The writer's own path runs here too, on the same evidence, writing to a directory the fixture owns.
    It writes both files before it prints, so a fault in the confirmation line leaves the published
    documents correct and the entry point returning a failure - which is a fault only running the entry
    point shows, and one this line shows without a third pass over the stack.

    The guide's glossary is checked in the direction a command can check: a term defined in the table
    and never used in the body is a dead entry, and it fails here rather than sitting unread.
    """
    from reporting import guide as guide_module
    from reporting import memo as memo_module

    root = Path(__file__).resolve().parents[1]
    evidence = memo_module.collect()
    assert (root / "reporting" / "findings-memo.md").read_text() == memo_module.memo(evidence)
    assert (root / "reporting" / "decision-record.md").read_text() == memo_module.record(evidence)
    assert (root / "reporting" / "learning-guide.md").read_text() == guide_module.guide(evidence)
    written = memo_module.write(evidence, tmp_path / "memo.md", tmp_path / "record.md")
    assert (tmp_path / "memo.md").read_text() == memo_module.memo(evidence)
    assert written["record"].read_text() == memo_module.record(evidence)
    assert guide_module.write(evidence, tmp_path / "guide.md").read_text() == guide_module.guide(evidence)

    unused = [term for term, appears in guide_module.coverage(evidence) if not appears]
    assert not unused, f"glossary terms the guide defines and never uses: {unused}"


def test_the_figures_exist_and_draw_no_instrument_series():
    """The figures are tracked artifacts, so they are read off disk rather than regenerated here.

    Two checks, both about the licence posture rather than about the drawing. Every figure the guide
    references is present and carries text, because a figure whose labels were written as glyph
    outlines is a figure a reader cannot search or quote. And no figure names an instrument: a chart
    of a sleeve's price or level would publish a series the feed's terms do not permit, and it could
    not be drawn without naming what it plots. The construction-time refusal of a level *kind* is the
    other half of the same rule, and it lives in the module that would have to draw one.
    """
    from reporting import figures as figures_module
    from reporting import guide as guide_module

    root = Path(__file__).resolve().parents[1]
    referenced = [
        root / "reporting" / "figures" / name for name, _ in guide_module.FIGURES_USED.values()
    ]
    assert referenced, "the guide references no figures"
    for path in referenced:
        assert path.exists(), f"{path.name} is referenced by the guide and absent from the repository"
        text = path.read_text()
        assert len(text.split()) > 20, f"{path.name} carries no text"
        assert "<text" in text, f"{path.name} was written with glyph outlines rather than text"
        named = [ticker for ticker in universe.TICKERS if ticker in text]
        assert not named, f"{path.name} names an instrument, so it may be drawing a series: {named}"
    for kind in ("price", "level", "nav", "index_level"):
        assert kind not in figures_module.KINDS, f"{kind} is a drawable kind of quantity"


def test_the_readme_tables_reprint_the_declarations():
    """The two data tables on the front page are restatements, so they are read back against the code.

    A README that lists twenty runs and eleven sleeves by hand is the most drift-prone artifact in the
    repository: a cell added to the registry or a weight changed in the universe would leave the page a
    reader meets first describing a grid the package no longer runs. The tables are therefore parsed
    out of the file and compared, cell by cell, against the declarations they restate. The prose is not
    checked, because prose is not a claim of equality with anything.
    """
    from portfolio_workbench.compare import registry
    from portfolio_workbench.construct import constraints
    from portfolio_workbench.data import universe

    root = Path(__file__).resolve().parents[1]
    text = (root / "README.md").read_text()
    rows = {}
    for line in text.split("\n"):
        if not line.startswith("| "):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        # Only the rows about an identifier are data: a table header is not keyed, because two tables
        # on the page may share one ("Family" heads the objectives and the test families alike).
        if not cells or not (cells[0].startswith("`") and cells[0].endswith("`")):
            continue
        if cells[0] in rows:
            raise AssertionError(f"the README carries two rows keyed {cells[0]}")
        rows[cells[0]] = cells

    for spec in registry.RUNS:
        key = f"`{spec['id']}`"
        assert key in rows, f"{spec['id']} is declared in the registry and absent from the README"
        row = rows[key]
        assert row[1] == spec["stage"], key
        assert row[2] == f"`{spec['family']}`", key
        assert row[3] == spec["estimator"], key
        assert row[4] == (spec["mean"] or "default"), key
        assert row[5] == spec["protocol"], key
        assert row[6] == ("35%" if spec["cap"] == constraints.CAP else "uncapped"), key

    for ticker, (sleeve, currency) in universe.SLEEVES.items():
        key = f"`{ticker}`"
        assert key in rows, f"{ticker} is in the universe and absent from the README"
        row = rows[key]
        assert row[1] == f"`{sleeve}`", key
        assert row[2] == universe.GROUP[sleeve], key
        assert row[3] == f"{universe.POLICY_WEIGHTS[ticker]:.0%}", key
        assert row[4] == currency, key

    assert len(rows) >= len(registry.RUNS) + len(universe.SLEEVES)


def test_the_readme_diagrams_are_well_formed():
    """A mermaid block that names an undefined node renders as an error box on the page.

    The check is structural rather than semantic: every block opens and closes, every label's brackets
    balance, and every node an edge refers to is defined somewhere in the same block. That catches the
    failure a typo produces, which is the failure a reader would see, and it does not pretend to
    validate mermaid itself.
    """
    root = Path(__file__).resolve().parents[1]
    text = (root / "README.md").read_text()
    blocks = re.findall(r"```mermaid\n(.*?)```", text, re.S)
    assert blocks, "the README carries no diagram"

    for block in blocks:
        assert block.count("[") == block.count("]"), block[:80]
        assert block.count("{") == block.count("}"), block[:80]
        defined = set()
        for line in block.split("\n"):
            for node in re.finditer(r"(\w+)\s*[\[\{]", line):
                defined.add(node.group(1))
            for subgraph in re.finditer(r"subgraph\s+(\w+)", line):
                defined.add(subgraph.group(1))
        edges = 0
        for line in block.split("\n"):
            parts = re.split(r"\s*(?:-->|-.->|---)\s*", line)
            for previous, following in zip(parts, parts[1:]):
                source = re.match(r"\s*(\w+)", previous)
                # An edge label leads the target part (`Q1 -->|no| V1`), so it is stripped before the
                # node id is read rather than being mistaken for one.
                target = re.match(r"(\w+)", re.sub(r"^\s*\|[^|]*\|\s*", "", following))
                assert source, f"an edge in the README starts with no node: {line.strip()[:60]}"
                assert source.group(1) in defined, f"{source.group(1)} is used and never defined"
                assert target, f"an edge in the README ends with no node: {line.strip()[:60]}"
                assert target.group(1) in defined, f"{target.group(1)} is used and never defined"
                edges += 1
        assert edges, "a diagram in the README carries no edge"


def test_every_method_is_traced_to_a_cited_source():
    """Done criterion 6: the source map is checked over the code, not over a hand-kept list."""
    missing = source_map.unmapped()
    assert not missing, f"public functions with no source: {missing}"
    mapped = source_map.entries()
    assert sum(len(functions) for functions in mapped.values()) == len(source_map.public_functions())


def test_every_skills_claim_is_backed_by_a_module_that_exists_and_is_exercised():
    """Done criterion 8: a covered row names a module, and this fixture reaches it."""
    problems = skills_matrix.unbacked(Path(__file__))
    assert not problems, problems


def test_the_workbook_carries_the_two_blocks_and_the_provenance_block(stack, tmp_path):
    """The output surface is checked on what it writes rather than on what it prints."""
    from openpyxl import load_workbook

    target = workbook.write(stack["sheet"], stack["document"], tmp_path / "comparison.xlsx")
    written = load_workbook(target)
    blocks = table_module.block_tables(stack["sheet"])
    assert written.sheetnames == [workbook.SHEET_NAMES[entry["block"]] for entry in blocks]

    provenance = table_module.provenance(stack["sheet"], stack["document"])
    for entry, name in zip(blocks, written.sheetnames):
        labels = [str(row[0]) for row in written[name].iter_rows(values_only=True)]
        for key in provenance:
            assert key in labels, f"{name} does not carry {key}"
        for table in entry["tables"]:
            assert table["title"] in labels, f"{name} does not carry {table['title']}"
    # Every row of the pre-registered table is in the sheet, under the cell id the table prints.
    labels = [str(row[0]) for row in written["returns and risk"].iter_rows(values_only=True)]
    for row in blocks[0]["tables"][0]["rows"]:
        assert str(row[0]) in labels, f"{row[0]} is missing from the workbook"
